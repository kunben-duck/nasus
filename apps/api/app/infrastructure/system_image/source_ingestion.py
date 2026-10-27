from __future__ import annotations

import hashlib
import io
import os
from pathlib import Path
from typing import Any, Sequence
from urllib.parse import urlparse

from ...application.system_image.source_ports import (
    IngestedSource,
    MaterializedSource,
    SourceConnector,
    SourceSpec,
    SourceTextUnit,
    SourceType,
    SourceAcquisitionDenied,
    SourceAcquisitionFailed,
)
from ...application.system_image.system_image_models import RawAssetRecord
from ...domain.system_image.source_tracking import require_tracked_ref
from .git_source_connector import GitSourceConnector
from .object_storage_source_connector import ObjectStorageSourceConnector
from .source_snapshot_archive import SourceSnapshotArchive
from .source_tree_scanner import SourceTreeScanner
from ..storage import ObjectStorage


class SourceIngestionService:
    def __init__(
        self,
        *,
        max_files: int | None = None,
        max_bytes: int | None = None,
        max_single_file_bytes: int | None = None,
        max_inventory_entries: int | None = None,
        allowed_external_schemes: set[str] | None = None,
        allowed_local_roots: list[Path] | None = None,
        connectors: Sequence[SourceConnector] | None = None,
        object_storage: ObjectStorage | None = None,
        snapshot_cache_dir: Path | None = None,
    ) -> None:
        storage = object_storage or ObjectStorage()
        self.snapshots = SourceSnapshotArchive(
            storage,
            snapshot_cache_dir or Path(os.getenv("NASUS_SOURCE_SNAPSHOT_CACHE_DIR", ".nasus/source-snapshots")),
        )
        self.max_files = max_files if max_files is not None else self._env_int("NASUS_SOURCE_MAX_FILES", 5000)
        self.max_bytes = max_bytes if max_bytes is not None else self._env_int("NASUS_SOURCE_MAX_BYTES", 100 * 1024 * 1024)
        self.max_single_file_bytes = (
            max_single_file_bytes
            if max_single_file_bytes is not None
            else self._env_int("NASUS_SOURCE_MAX_SINGLE_FILE_BYTES", 10 * 1024 * 1024)
        )
        self.max_inventory_entries = (
            max_inventory_entries if max_inventory_entries is not None
            else self._env_int("NASUS_SOURCE_MAX_INVENTORY_ENTRIES", 20000)
        )
        self.allowed_external_schemes = allowed_external_schemes or self._env_csv_set(
            "NASUS_SOURCE_ALLOWED_URI_SCHEMES",
            {"git", "http", "https", "ssh", "s3", "local-object", "gs", "docs", "tests"},
        )
        self.allowed_local_roots = allowed_local_roots if allowed_local_roots is not None else self._env_path_list(
            "NASUS_SOURCE_ALLOWED_LOCAL_ROOTS"
        )
        self.connectors = (
            tuple(connectors)
            if connectors is not None
            else (GitSourceConnector(), ObjectStorageSourceConnector(storage))
        )
        self._materializations: dict[tuple[str, str, str, str | None, str | None], MaterializedSource] = {}

    def normalize_specs(self, raw_specs: Any) -> list[SourceSpec]:
        if not isinstance(raw_specs, list):
            return []

        specs: list[SourceSpec] = []
        for item in raw_specs:
            if not isinstance(item, dict):
                continue
            source_type = item.get("source_type") or item.get("type")
            source_uri = item.get("source_uri") or item.get("uri") or item.get("path")
            if source_type not in {"code", "us_doc", "test_asset"} or not isinstance(source_uri, str):
                continue
            label = item.get("label")
            credential_ref = item.get("credential_ref")
            specs.append(
                SourceSpec(
                    source_type=source_type,
                    source_uri=source_uri.strip(),
                    label=label.strip() if isinstance(label, str) and label.strip() else None,
                    credential_ref=credential_ref.strip()
                    if isinstance(credential_ref, str) and credential_ref.strip()
                    else None,
                    tracked_ref=item.get("tracked_ref"),
                )
            )
        return specs

    def ingest(self, source: RawAssetRecord) -> IngestedSource:
        require_tracked_ref(source.source_type, source.tracked_ref)
        path = self._local_path(source.source_uri)
        if path is not None and source.tracked_ref is None:
            return self._ingest_local_path(source, path)
        if path is None:
            self._assert_external_scheme_allowed(source.source_uri)
        else:
            self._assert_local_path_allowed(path)
        materialized = self._materialize_external_source(source, refresh=True)
        if materialized is not None:
            ingested = self._ingest_local_path(
                source,
                materialized.local_path,
                enforce_allowed_root=False,
                external_revision=materialized.revision,
            )
            return IngestedSource(
                content_hash=ingested.content_hash,
                content_ref=ingested.content_ref,
                evidence_refs=[*materialized.evidence_refs, *ingested.evidence_refs],
                file_count=ingested.file_count,
                byte_count=ingested.byte_count,
                snapshot=ingested.snapshot,
                inventory=ingested.inventory,
            )
        if source.source_type == "code":
            raise OSError(f"no source connector can materialize remote code source: {source.source_uri}")
        return self._ingest_external_reference(source)

    def extract_text_units(self, source: RawAssetRecord) -> list[SourceTextUnit]:
        require_tracked_ref(source.source_type, source.tracked_ref)
        archived = self.snapshots.owns_ref(source.content_ref)
        path = (
            self.snapshots.materialize(
                source.content_ref, project_id=source.project_id, source_id=source.id,
                content_hash=source.content_hash, max_files=self.max_files,
                max_bytes=self.max_bytes, max_file_bytes=self.max_single_file_bytes,
            )
            if archived else self._local_path(source.source_uri) if source.tracked_ref is None else None
        )
        if path is None:
            local_path = self._local_path(source.source_uri)
            if local_path is None:
                self._assert_external_scheme_allowed(source.source_uri)
            else:
                self._assert_local_path_allowed(local_path)
            materialized = self._materialize_external_source(source, refresh=False)
            if materialized is not None:
                path = materialized.local_path
            else:
                if source.tracked_ref is not None:
                    raise OSError("no Git connector can materialize the explicit tracked_ref")
                digest = hashlib.sha256(source.source_uri.encode("utf-8")).hexdigest()
                return [
                    SourceTextUnit(
                        relative_path=f"{source.source_type}:external-reference",
                        text="\n".join([
                            f"source_type: {source.source_type}",
                            f"source_uri: {source.source_uri}",
                            f"content_hash: {source.content_hash}",
                            "connector_status: external connector unavailable",
                        ]),
                        byte_count=0,
                        content_hash=f"ref-sha256:{digest}",
                    )
                ]

        if not archived and source.tracked_ref is None and self._local_path(source.source_uri) is not None:
            self._assert_local_path_allowed(path)
        if not path.exists():
            raise FileNotFoundError(f"source path does not exist: {path}")

        files = self._ingestible_files(path)
        units: list[SourceTextUnit] = []
        byte_count = 0
        source_root = str(path) if path.is_dir() else None
        for file_path in files:
            relative = file_path.name if path.is_file() else str(file_path.relative_to(path))
            file_size = file_path.stat().st_size
            if file_size > self.max_single_file_bytes:
                raise OSError(
                    f"source file exceeds max single file bytes {self.max_single_file_bytes}: {relative}"
                )
            if byte_count + file_size > self.max_bytes:
                raise OSError(f"source exceeds max total bytes {self.max_bytes}: {source.source_type}")
            data = file_path.read_bytes()
            byte_count += file_size
            text = self._decode_file(file_path, data)
            if not text.strip():
                continue
            units.append(
                SourceTextUnit(
                    relative_path=relative,
                    text=text,
                    byte_count=file_size,
                    content_hash=f"sha256:{hashlib.sha256(data).hexdigest()}",
                    source_root=source_root,
                    text_is_lossless=text.encode("utf-8") == data,
                )
            )
        return units

    def _ingest_local_path(
        self,
        source: RawAssetRecord,
        path: Path,
        *,
        enforce_allowed_root: bool = True,
        external_revision: str | None = None,
    ) -> IngestedSource:
        if enforce_allowed_root:
            self._assert_local_path_allowed(path)
        scan = SourceTreeScanner(
            excluded_names=self.ignored_parts, max_files=self.max_files,
            max_bytes=self.max_bytes, max_file_bytes=self.max_single_file_bytes,
            max_entries=self.max_inventory_entries,
        ).scan(path)
        digest = hashlib.sha256()
        for relative, data in scan.files:
            digest.update(relative.encode("utf-8"))
            digest.update(b"\0")
            digest.update(data)
        content_hash = f"sha256:{digest.hexdigest()}"
        inventory = self.snapshots.store_inventory(
            scan.inventory, project_id=source.project_id, source_id=source.id,
            external_revision=external_revision or (content_hash if scan.inventory.complete else None),
        )
        if not scan.inventory.complete:
            denied = any(item.reason in {"permission_denied", "symbolic_link"} for item in scan.inventory.entries)
            raise (SourceAcquisitionDenied if denied else SourceAcquisitionFailed)(inventory)
        try:
            snapshot = self.snapshots.capture(
                project_id=source.project_id, source_id=source.id, files=list(scan.files),
                content_hash=content_hash, external_revision=external_revision,
                excluded_directories=sorted(self.ignored_parts), inventory=scan.inventory,
            )
        except (OSError, RuntimeError, ValueError) as exc:
            raise SourceAcquisitionFailed(inventory) from exc
        return IngestedSource(
            content_hash=content_hash,
            content_ref=snapshot.manifest_ref,
            evidence_refs=self._typed_evidence_refs(
                source.source_type, [f"file:{relative}" for relative, _ in scan.files],
                inventory.file_count, inventory.byte_count,
            ),
            file_count=inventory.file_count,
            byte_count=inventory.byte_count,
            snapshot=snapshot,
            inventory=inventory,
        )

    def _ingest_external_reference(self, source: RawAssetRecord) -> IngestedSource:
        digest = hashlib.sha256(source.source_uri.encode("utf-8")).hexdigest()
        return IngestedSource(
            content_hash=f"ref-sha256:{digest}",
            content_ref=f"nasus://raw/{source.project_id}/{source.source_type}/{digest[:16]}",
            evidence_refs=self._typed_evidence_refs(source.source_type, [f"uri:{source.source_uri}"], 0, 0),
            file_count=0,
            byte_count=0,
        )

    def _materialize_external_source(
        self,
        source: RawAssetRecord,
        *,
        refresh: bool,
    ) -> MaterializedSource | None:
        pinned_revision = (
            GitSourceConnector._pinned_revision(source)
            if source.source_type == "code"
            else None
        )
        binding_key = (source.project_id, source.id, source.source_uri, source.tracked_ref)
        cache_key = (*binding_key, pinned_revision)
        if not refresh and cache_key in self._materializations:
            cached = self._materializations[cache_key]
            if cached.connector_kind != "git":
                return cached
        connector = next((item for item in self.connectors if item.supports(source)
            and (source.tracked_ref is None or item.connector_kind == "git")), None)
        if connector is None:
            return None
        materialized = connector.materialize(source, refresh=refresh)
        if materialized.connector_kind == "git" or source.source_type == "code":
            self._materializations[(*binding_key, materialized.revision)] = materialized
            if refresh or pinned_revision is None:
                self._materializations[(*binding_key, None)] = materialized
        else:
            self._materializations[cache_key] = materialized
        return materialized

    def _ingestible_files(self, path: Path) -> list[Path]:
        candidates = [path] if path.is_file() else sorted(path.rglob("*"))
        files: list[Path] = []
        for item in candidates:
            if self._is_ignored_path(Path(item.name) if path.is_file() else item.relative_to(path)):
                continue
            if item.is_symlink():
                raise PermissionError("source contains a symbolic link; bind a regular-file snapshot")
            if item.is_file():
                files.append(item)
        if not files:
            raise OSError(f"source path contains no ingestible files: {path}")
        if len(files) > self.max_files:
            raise OSError(f"source exceeds max file count {self.max_files}: {len(files)} files")
        return files

    @staticmethod
    def _typed_evidence_refs(source_type: SourceType, refs: list[str], file_count: int, byte_count: int) -> list[str]:
        if source_type == "code":
            base = ["source:code", "parser:file-fingerprint"]
        elif source_type == "us_doc":
            base = ["source:historical-us", "parser:document-fingerprint"]
        else:
            base = ["source:test-assets", "parser:test-fingerprint"]
        return [*base, f"files:{file_count}", f"bytes:{byte_count}", *refs[:20]]

    @staticmethod
    def _local_path(source_uri: str) -> Path | None:
        parsed = urlparse(source_uri)
        if parsed.scheme == "file":
            return Path(parsed.path).expanduser().resolve()
        if parsed.scheme:
            return None
        return Path(source_uri).expanduser().resolve()

    def _assert_external_scheme_allowed(self, source_uri: str) -> None:
        parsed = urlparse(source_uri)
        if parsed.scheme not in self.allowed_external_schemes:
            raise PermissionError(f"source URI scheme is not allowed: {parsed.scheme or 'unknown'}")

    def _assert_local_path_allowed(self, path: Path) -> None:
        if not self.allowed_local_roots:
            return
        resolved = path.resolve()
        for root in self.allowed_local_roots:
            if resolved == root or root in resolved.parents:
                return
        raise PermissionError(f"source local path is outside allowed roots: {resolved}")

    ignored_parts = frozenset({
            ".git",
            ".nasus-complete",
            "node_modules",
            "__pycache__",
            ".venv",
            "dist",
            "build",
    })

    @classmethod
    def _is_ignored_path(cls, path: Path) -> bool:
        return any(part in cls.ignored_parts for part in path.parts)

    @classmethod
    def _decode_file(cls, file_path: Path, data: bytes) -> str:
        suffix = file_path.suffix.lower()
        try:
            if suffix == ".pdf":
                from pypdf import PdfReader

                reader = PdfReader(io.BytesIO(data))
                return "\n\n".join(
                    f"[page {index}]\n{page.extract_text() or ''}"
                    for index, page in enumerate(reader.pages, start=1)
                )
            if suffix == ".docx":
                from docx import Document

                document = Document(io.BytesIO(data))
                paragraphs = [item.text for item in document.paragraphs if item.text.strip()]
                tables = [
                    "\n".join(
                        "\t".join(cell.text for cell in row.cells)
                        for row in table.rows
                    )
                    for table in document.tables
                ]
                return "\n\n".join([*paragraphs, *tables])
            if suffix in {".xlsx", ".xlsm"}:
                from openpyxl import load_workbook

                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                sheets: list[str] = []
                try:
                    for sheet in workbook.worksheets:
                        rows = [
                            "\t".join("" if value is None else str(value) for value in row)
                            for row in sheet.iter_rows(values_only=True)
                        ]
                        sheets.append(f"[sheet {sheet.title}]\n" + "\n".join(rows))
                finally:
                    workbook.close()
                return "\n\n".join(sheets)
        except Exception as exc:
            raise OSError(f"failed to extract text from {file_path.name}: {exc}") from exc
        return cls._decode_text(data)

    @staticmethod
    def _decode_text(data: bytes) -> str:
        if b"\0" in data[:4096]:
            return ""
        return data.decode("utf-8", errors="replace")

    @staticmethod
    def _env_int(name: str, default: int) -> int:
        raw = os.getenv(name)
        if not raw:
            return default
        try:
            return max(1, int(raw))
        except ValueError:
            return default

    @staticmethod
    def _env_csv_set(name: str, default: set[str]) -> set[str]:
        raw = os.getenv(name)
        if not raw:
            return set(default)
        values = {item.strip().lower() for item in raw.split(",") if item.strip()}
        return values or set(default)

    @staticmethod
    def _env_path_list(name: str) -> list[Path]:
        raw = os.getenv(name)
        if not raw:
            return []
        return [Path(item).expanduser().resolve() for item in raw.split(os.pathsep) if item.strip()]


__all__ = [
    "IngestedSource",
    "SourceIngestionService",
    "SourceSpec",
    "SourceTextUnit",
    "SourceType",
]
