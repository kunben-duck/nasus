from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook

from apps.api.app.application.system_image.source_uploads import (
    SourceUploadApplicationService,
    SourceUploadFile,
    SourceUploadResult,
    UploadSourceType,
)
from apps.api.app.application.system_image.system_image_models import RawAssetRecord
from apps.api.app.infrastructure.storage.object_storage import (
    ObjectStorage,
    ObjectStorageConfig,
)
from apps.api.app.infrastructure.system_image.object_storage_source_connector import (
    ObjectStorageSourceConnector,
    ObjectStorageSourceConnectorError,
)
from apps.api.app.infrastructure.system_image.source_ingestion import SourceIngestionService
from apps.api.app.infrastructure.system_image.source_upload_storage import (
    ObjectStorageSourceUploadStorage,
)


class FakeWorkspace:
    def __init__(self, project_ids: set[str]) -> None:
        self.project_ids = project_ids
        self.access_checks: list[str] = []

    def has_project(self, project_id: str) -> bool:
        return project_id in self.project_ids

    def require_project_access(self, project_id: str) -> None:
        self.access_checks.append(project_id)


@dataclass
class RecordingStorage:
    result: SourceUploadResult
    uploads: list[tuple[str, UploadSourceType, Sequence[SourceUploadFile]]]

    def put_source_bundle(
        self,
        project_id: str,
        source_type: UploadSourceType,
        files: Sequence[SourceUploadFile],
    ) -> SourceUploadResult:
        self.uploads.append((project_id, source_type, files))
        return self.result


def local_object_storage(tmp_path: Path) -> ObjectStorage:
    return ObjectStorage(
        ObjectStorageConfig(
            endpoint=None,
            bucket="source-test",
            access_key=None,
            secret_key=None,
            region="us-east-1",
            local_dir=tmp_path / "objects",
        )
    )


def docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def xlsx_bytes(text: str) -> bytes:
    workbook = Workbook()
    workbook.active.title = "Cases"
    workbook.active.append(["case", "expected"])
    workbook.active.append([text, "pass"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_source_upload_application_validates_project_scope_paths_and_limits() -> None:
    workspace = FakeWorkspace({"proj_1"})
    result = SourceUploadResult(
        source_type="us_doc",
        source_uri="local-object://test/source.zip",
        content_hash="sha256:test",
        file_count=1,
        byte_count=3,
        evidence_refs=[],
    )
    storage = RecordingStorage(result=result, uploads=[])
    service = SourceUploadApplicationService(
        workspace,  # type: ignore[arg-type]
        storage,
        max_files=2,
        max_bytes=8,
        max_single_file_bytes=5,
    )

    with pytest.raises(KeyError):
        service.authorize_project("missing")
    service.authorize_project("proj_1")
    assert workspace.access_checks[-1] == "proj_1"
    with pytest.raises(KeyError):
        service.upload("missing", "us_doc", [SourceUploadFile("US.md", "text/markdown", b"US")])
    with pytest.raises(ValueError, match="source_type"):
        service.upload("proj_1", "code", [SourceUploadFile("code.py", "text/plain", b"x")])
    with pytest.raises(ValueError, match="invalid|unsafe path"):
        service.upload("proj_1", "us_doc", [SourceUploadFile("../US.md", "text/markdown", b"US")])
    with pytest.raises(ValueError, match="single file"):
        service.upload("proj_1", "us_doc", [SourceUploadFile("US.md", "text/markdown", b"123456")])

    uploaded = service.upload(
        "proj_1",
        "us_doc",
        [SourceUploadFile("stories/US.md", "text/markdown", b"US1")],
    )
    assert uploaded == result
    assert workspace.access_checks[-1] == "proj_1"
    assert storage.uploads[-1][2][0].filename == "stories/US.md"


def test_object_storage_upload_round_trip_extracts_office_and_text_sources(tmp_path: Path) -> None:
    object_storage = local_object_storage(tmp_path)
    workspace = FakeWorkspace({"proj_upload"})
    service = SourceUploadApplicationService(
        workspace,  # type: ignore[arg-type]
        ObjectStorageSourceUploadStorage(object_storage),
    )
    uploaded = service.upload(
        "proj_upload",
        "us_doc",
        [
            SourceUploadFile("US-201.md", "text/markdown", b"# US-201\nSaved card checkout"),
            SourceUploadFile(
                "acceptance/criteria.docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                docx_bytes("Checkout preserves the selected card"),
            ),
            SourceUploadFile(
                "acceptance/matrix.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                xlsx_bytes("Saved card checkout"),
            ),
        ],
    )

    assert uploaded.source_uri.startswith("local-object://source-test/source-uploads/")
    assert uploaded.file_count == 3
    assert uploaded.content_hash.startswith("sha256:")

    source = RawAssetRecord(
        id="raw_upload",
        project_id="proj_upload",
        source_type="us_doc",
        source_uri=uploaded.source_uri,
    )
    connector = ObjectStorageSourceConnector(
        object_storage,
        cache_dir=tmp_path / "cache",
    )
    ingestion = SourceIngestionService(connectors=(connector,))
    ingested = ingestion.ingest(source)
    units = ingestion.extract_text_units(source)
    combined_text = "\n".join(item.text for item in units)

    assert ingested.file_count == 3
    assert "connector:object-storage-bundle" in ingested.evidence_refs
    assert "Saved card checkout" in combined_text
    assert "Checkout preserves the selected card" in combined_text
    assert "[sheet Cases]" in combined_text


def test_object_storage_connector_rejects_zip_path_traversal(tmp_path: Path) -> None:
    object_storage = local_object_storage(tmp_path)
    archive_body = io.BytesIO()
    with ZipFile(archive_body, "w") as archive:
        archive.writestr("../escape.txt", "unsafe")
    stored = object_storage.put_bytes(
        "source-uploads/proj/us_doc/unsafe.zip",
        archive_body.getvalue(),
        content_type="application/zip",
    )
    source = RawAssetRecord(
        id="raw_unsafe",
        project_id="proj_upload",
        source_type="us_doc",
        source_uri=stored.storage_ref,
    )

    connector = ObjectStorageSourceConnector(
        object_storage,
        cache_dir=tmp_path / "cache",
    )
    with pytest.raises(ObjectStorageSourceConnectorError, match="unsafe path"):
        connector.materialize(source, refresh=True)
